"""Turn an uploaded resume (PDF / DOCX / TXT) into plain text."""
import io
import re
from pathlib import Path


class ParseError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ParseError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    text = _clean(text)
    if len(text) < 100:
        raise ParseError(
            "Very little text could be extracted. If this is a scanned/image PDF, paste the resume text instead."
        )
    return text


def _pdf(data: bytes) -> str:
    import pdfplumber

    pages, links = [], []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
                links += [h.get("uri") for h in (page.hyperlinks or []) if h.get("uri")]
    except Exception as e:
        raise ParseError(f"Could not read this PDF: {e}") from e
    return "\n".join(pages) + _links_block(links)


def _docx(data: bytes) -> str:
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not read this DOCX: {e}") from e
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
            lines.append(" | ".join(cells))
    links = [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    return "\n".join(lines) + _links_block(links)


def _links_block(links: list[str]) -> str:
    # Links hidden behind icons/text (e.g. a LinkedIn icon) are lost in plain text, so list them.
    links = [l for l in dict.fromkeys(links) if not l.startswith("mailto:")]
    return "\n\nLinks found in document:\n" + "\n".join(links) if links else ""


def _clean(text: str) -> str:
    text = text.replace("\r", "").replace(" ", " ")
    text = re.sub(r"[●▪■•➢]\s*", "• ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
